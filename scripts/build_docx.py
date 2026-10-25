"""Render report/paper.md into the Apart submission template.

We build ON TOP of the supplied template rather than generating a fresh
document, so Apart's own Title/Heading/Normal styles and page setup are kept
and the submission looks like the template they asked for. The template's
guidance paragraphs are deleted first.

This is a deliberately small markdown subset -- headings, paragraphs with
bold/italic/code spans, bullet and numbered lists, pipe tables, fenced code
blocks, and image directives -- because the input is one known document, not
arbitrary markdown.

    python scripts/build_docx.py            # build
    python scripts/build_docx.py --stats    # report length only

Figures are inserted where the source contains a line of the form
    {{FIGURE: path | caption}}
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

TEMPLATE = Path("report/template/apart_template.docx")
SOURCE = Path("report/paper.md")
OUT = Path("report/Swayambhu_Apart_submission.docx")

INK_MUTED = RGBColor(0x52, 0x51, 0x4E)


def has_style(doc, name: str) -> bool:
    try:
        doc.styles[name]
        return True
    except KeyError:
        return False


def set_borders(table) -> None:
    """The template ships no Table Grid style, so draw borders directly."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "BFBFBF")
        borders.append(el)
    tblPr.append(borders)


def clear_body(doc) -> None:
    """Remove every paragraph and table the template ships with."""
    body = doc.element.body
    for child in list(body):
        if child.tag.endswith(("}p", "}tbl")):
            body.remove(child)


# --- inline formatting -----------------------------------------------------

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_runs(par, text: str, size: int | None = None) -> None:
    """Write text into a paragraph, honouring **bold**, *italic* and `code`."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            r = par.add_run(piece[2:-2]); r.bold = True
        elif piece.startswith("*") and piece.endswith("*"):
            r = par.add_run(piece[1:-1]); r.italic = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1]); r.font.name = "Consolas"
            r.font.size = Pt((size or 11) - 1)
        else:
            r = par.add_run(piece)
        if size:
            r.font.size = Pt(size)


def add_table(doc, rows: list[list[str]]) -> None:
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=len(rows), cols=len(header))
    if has_style(doc, "Table Grid"):
        t.style = "Table Grid"
    else:
        set_borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, cell in enumerate(header):
        p = t.cell(0, j).paragraphs[0]
        add_runs(p, cell or "", size=9)
        for r in p.runs:
            r.bold = True
    for i, row in enumerate(body, start=1):
        for j, cell in enumerate(row[: len(header)]):
            add_runs(t.cell(i, j).paragraphs[0], cell or "", size=9)
    doc.add_paragraph()


def add_figure(doc, path: str, caption: str, n: int) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(6.0))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(f"Figure {n}. ")
    r.bold = True; r.font.size = Pt(9); r.font.color.rgb = INK_MUTED
    r2 = cap.add_run(caption)
    r2.font.size = Pt(9); r2.font.color.rgb = INK_MUTED


# --- document assembly -----------------------------------------------------

def build(stats_only: bool = False) -> None:
    src = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = docx.Document(str(TEMPLATE))
    clear_body(doc)

    i = 0
    fig_n = 0
    pending_table: list[list[str]] = []

    def flush_table():
        nonlocal pending_table
        if pending_table:
            add_table(doc, pending_table)
            pending_table = []

    while i < len(src):
        line = src[i].rstrip()

        # pipe table rows accumulate until a non-table line
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if not all(set(c) <= set("-: ") for c in cells):   # skip separator
                pending_table.append(cells)
            i += 1
            continue
        flush_table()

        if not line:
            i += 1
            continue

        if line.startswith("{{FIGURE:"):
            inner = line[len("{{FIGURE:"):].rstrip("}").strip()
            path, _, caption = inner.partition("|")
            fig_n += 1
            if not stats_only:
                add_figure(doc, path.strip(), caption.strip(), fig_n)
            i += 1
            continue

        if line.startswith("```"):                              # fenced block
            i += 1
            buf = []
            while i < len(src) and not src[i].startswith("```"):
                buf.append(src[i]); i += 1
            i += 1
            for b in buf:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Inches(0.25)
                r = p.add_run(b)
                r.font.name = "Consolas"; r.font.size = Pt(8.5)
            doc.add_paragraph()
            continue

        if line.startswith("# "):
            doc.add_paragraph(line[2:], style="Title"); i += 1; continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 3"); i += 1; continue
        if line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 2"); i += 1; continue
        if line.startswith("---"):
            i += 1; continue

        m = re.match(r"^(\d+)\.\s+(.*)", line)
        if m:
            # the template defines no list styles, so indent manually rather
            # than failing on a style that is not there
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, f"{m.group(1)}.  {m.group(2)}")
            i += 1; continue
        if line.startswith("- "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, "•  " + line[2:])
            i += 1; continue

        p = doc.add_paragraph()
        add_runs(p, line)
        i += 1

    flush_table()

    if stats_only:
        words = sum(len(p.text.split()) for p in doc.paragraphs)
        print(f"paragraphs {len(doc.paragraphs)}  tables {len(doc.tables)}  "
              f"figures {fig_n}  words {words}")
        return

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"wrote {OUT}  ({len(doc.paragraphs)} paragraphs, "
          f"{len(doc.tables)} tables, {fig_n} figures)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--stats", action="store_true")
    build(ap.parse_args().stats)
